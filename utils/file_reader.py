import io


def read_file(uploaded_file):
    """
    Read text from an uploaded PDF, DOCX, or TXT file.

    Supported formats:
    - PDF
    - DOCX
    - TXT
    """

    if uploaded_file is None:
        raise ValueError("No file was uploaded.")

    file_name = uploaded_file.name.lower()

    # =========================================================
    # PDF
    # =========================================================

    if file_name.endswith(".pdf"):

        try:
            from PyPDF2 import PdfReader

            file_bytes = uploaded_file.getvalue()

            reader = PdfReader(
                io.BytesIO(file_bytes)
            )

            text_parts = []

            for page in reader.pages:

                page_text = page.extract_text()

                if page_text:
                    text_parts.append(page_text)

            return "\n".join(text_parts).strip()

        except ImportError:

            raise RuntimeError(
                "PDF support requires PyPDF2. "
                "Please run: pip install PyPDF2"
            )

        except Exception as exc:

            raise RuntimeError(
                f"Could not read the PDF file: {exc}"
            ) from exc

    # =========================================================
    # DOCX
    # =========================================================

    elif file_name.endswith(".docx"):

        try:
            from docx import Document

            file_bytes = uploaded_file.getvalue()

            document = Document(
                io.BytesIO(file_bytes)
            )

            text_parts = []

            # Read normal paragraphs
            for paragraph in document.paragraphs:

                text = paragraph.text.strip()

                if text:
                    text_parts.append(text)

            # Read tables too
            for table in document.tables:

                for row in table.rows:

                    row_text = []

                    for cell in row.cells:

                        cell_text = cell.text.strip()

                        if cell_text:
                            row_text.append(cell_text)

                    if row_text:
                        text_parts.append(
                            " | ".join(row_text)
                        )

            return "\n".join(text_parts).strip()

        except ImportError:

            raise RuntimeError(
                "DOCX support requires python-docx. "
                "Please run: pip install python-docx"
            )

        except Exception as exc:

            raise RuntimeError(
                f"Could not read the DOCX file: {exc}"
            ) from exc

    # =========================================================
    # TXT
    # =========================================================

    elif file_name.endswith(".txt"):

        try:

            file_bytes = uploaded_file.getvalue()

            return file_bytes.decode(
                "utf-8",
                errors="replace"
            ).strip()

        except Exception as exc:

            raise RuntimeError(
                f"Could not read the TXT file: {exc}"
            ) from exc

    # =========================================================
    # UNSUPPORTED FORMAT
    # =========================================================

    else:

        raise ValueError(
            "Unsupported file format. "
            "Please upload a PDF, DOCX, or TXT file."
        )